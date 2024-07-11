import time
import requests
import json
from openpyxl import Workbook
import pandas as pd
import pymysql
import traceback
from datetime import datetime
from py_mini_racer import py_mini_racer
from dotenv import load_dotenv
import os
from fake_useragent import UserAgent
from send_data import send_data 

# Load environment variables
load_dotenv()

# Setup constants and configurations
BASE_URL = os.environ.get('BASE_URL')
DB_CONFIG = {
    'host': os.environ.get('HOST'),
    'user': os.environ.get('DB_USER'),
    'password': os.environ.get('DB_PASSWORD'),
    'db': os.environ.get('DB_TABLE'),
    'charset': os.environ.get('CHARSET')
}
SAVE_PATH = os.environ.get('SAVE_PATH')
EXCEL_PATH = os.path.join(SAVE_PATH, 'kream_test_data.xlsx')
CSV_PATH = os.path.join(SAVE_PATH, 'kream_test_data.csv')
LOG_FILE = 'crawling(bot_3).log'

# Initialize log file
def init_log(file_path):
    with open(file_path, 'a', encoding='euc-kr') as f:
        f.write(f'start bot3 at : {datetime.now()}\n')

# Write log
def write_log(file_path, message):
    with open(file_path, 'a', encoding='euc-kr') as f:
        f.write(message + '\n')

# Fetch the last crawled product ID
def get_last_crawled_p_id(cur):
    cur.execute("SELECT * FROM config WHERE cfg_key = 'crawled_p_id'")
    result = cur.fetchone()
    if result is None:
        cur.execute("INSERT INTO config(cfg_key, cfg_value) VALUES('crawled_p_id', 0)")
        return 0
    return int(result['cfg_value'])

# Fetch products to crawl
def fetch_products(cur, last_crawled_p_id, p_count=80, limit=500):
    cur.execute(f"SELECT * FROM product WHERE p_id > {last_crawled_p_id} AND p_count >= {p_count} LIMIT {limit}")
    return cur.fetchall()

# Update the last crawled product ID
def update_last_crawled_p_id(cur, p_id):
    cur.execute(f"UPDATE config SET cfg_value = {p_id} WHERE cfg_key = 'crawled_p_id'")

# Extract script content from HTML
def extract_script_content(html_content):
    script_start = html_content.find("window.__NUXT__=")
    script_end = html_content.find("</script>", script_start)
    if script_start == -1 or script_end == -1:
        return None
    return html_content[script_start:script_end]

# Parse nuxt data using py_mini_racer
def parse_nuxt_data(script_content):
    ctx = py_mini_racer.MiniRacer()
    ctx.eval("window={};")
    ctx.eval(script_content)
    try:
        return ctx.eval("JSON.stringify(window.__NUXT__.state.product.response.sales_options)")
    except:
        return ctx.eval("JSON.stringify(window.__NUXT__.state.product.response.sale_info.details.options)")

# Insert or update product prices
def upsert_product_prices(cur, product_id, items):
    for item in items:
        if 'is_show_option_price' in item and not item['is_show_option_price']:
            continue

        price = item.get('price', item.get('lowest_normal', 0))
        fast_price = item.get('lowest_100', 0)
        option = item.get('key', item.get('option', ''))

        cur.execute(
            f"""
            INSERT INTO price_by_option_of_product(p_id, pboop_option, pboop_price, pboop_quick_price) 
            VALUES (%s, %s, %s, %s)
            ON DUPLICATE KEY UPDATE 
                pboop_price = VALUES(pboop_price), 
                pboop_quick_price = VALUES(pboop_quick_price), 
                pboop_datetime = %s
            """,
            (product_id, option, price, fast_price, datetime.now())
        )

# Process each product
def process_product(session, product, cur, ua, headers):
    response = session.get(product['p_url'], headers=headers)
    if response.status_code != 200:
        return False, product['p_id']
    
    script_content = extract_script_content(response.text)
    if script_content is None:
        write_log(LOG_FILE, "script not found")
        return False, product['p_id']
    
    nuxt_data = parse_nuxt_data(script_content)
    if nuxt_data is None:
        write_log(LOG_FILE, "failed to parse nuxt data")
        return False, product['p_id']
    
    upsert_product_prices(cur, product['p_id'], json.loads(nuxt_data))
    write_log(LOG_FILE, f"success. p_id : {product['p_id']}")
    return True, product['p_id']

# Main function to perform crawling
def main():
    ua = UserAgent()
    ua.browsers.remove('chrome')

    init_log(LOG_FILE)
    session = requests.Session()

    try:
        conn = pymysql.connect(**DB_CONFIG)
        cur = conn.cursor(pymysql.cursors.DictCursor)

        last_crawled_p_id = get_last_crawled_p_id(cur)
        products = fetch_products(cur, last_crawled_p_id)
        write_log(LOG_FILE, f"bot3 : loaded products {len(products)}")

        headers = {'user-agent': ua.random}
        for i, product in enumerate(products):
            if i % 400 == 0:
                headers = {'user-agent': ua.random}
                print('change user-agent')

            success, p_id = process_product(session, product, cur, ua, headers)
            if not success:
                update_last_crawled_p_id(cur, p_id - 1)
                break

            update_last_crawled_p_id(cur, p_id)
            write_log(LOG_FILE, f"p_id : {p_id} end at :{datetime.now()}")
            conn.commit()

        # Write results to Excel and CSV
        cur.execute(
            "SELECT p_name, p_url, pboop_option, pboop_price, pboop_quick_price, pboop_datetime "
            "FROM price_by_option_of_product "
            "JOIN product ON price_by_option_of_product.p_id = product.p_id"
        )
        rows = cur.fetchall()
        result = pd.DataFrame(rows)
        
        if not os.path.exists(EXCEL_PATH):
            wb = Workbook()
            ws = wb.active
            ws.title = "result_sheet"
            wb.save(EXCEL_PATH)

        with pd.ExcelWriter(EXCEL_PATH, engine='openpyxl', mode='a', if_sheet_exists='replace') as writer:
            if "result_sheet" in writer.book.sheetnames:
                idx = writer.book.sheetnames.index("result_sheet")
                std = writer.book.worksheets[idx]
                writer.book.remove(std)

            result.to_excel(writer, sheet_name='result_sheet', index=False)

        result.to_csv(CSV_PATH, index=False)

        send_data()

    except Exception:
        err = traceback.format_exc()
        write_log(LOG_FILE, f'error bot3 : {err}')

    finally:
        write_log(LOG_FILE, f'ended bot3 at : {datetime.now()}')
        session.close()
        conn.close()

if __name__ == "__main__":
    main()